from __future__ import annotations

from pathlib import Path
from typing import Iterable, List, Tuple

from PIL import Image
from pptx import Presentation


def list_images(image_dir: Path, recursive: bool, extensions: Iterable[str]) -> List[Path]:
    exts = {e.lower() for e in extensions}
    globber = image_dir.rglob if recursive else image_dir.glob
    files = [p for p in globber("*") if p.is_file() and p.suffix.lower() in exts]
    return sorted(files)


def extract_images_from_pptx(pptx_path: Path, output_dir: Path) -> List[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(pptx_path))
    written: List[Path] = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        image_idx = 0
        for shape in slide.shapes:
            if shape.shape_type != 13:
                continue
            image_idx += 1
            img = shape.image
            ext = img.ext if img.ext else "png"
            out_path = output_dir / f"slide_{slide_idx:02d}_img_{image_idx:02d}.{ext}"
            out_path.write_bytes(img.blob)
            written.append(out_path)
    return written


def load_rgb_image(image_path: Path):
    with Image.open(image_path) as img:
        return img.convert("RGB")


def ensure_dirs(root: Path) -> Tuple[Path, Path]:
    records = root / "records"
    artifacts = root / "artifacts"
    records.mkdir(parents=True, exist_ok=True)
    artifacts.mkdir(parents=True, exist_ok=True)
    return records, artifacts
